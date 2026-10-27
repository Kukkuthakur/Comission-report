import io
import re
import urllib.request
from datetime import date

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------
# Period label helpers
# ---------------------------------------------------------------
def default_period_text(today=None):
    """Return '1 Mon - D Mon YYYY' for the current month up to today."""
    today = today or date.today()
    first = today.replace(day=1)
    return f"1 {first:%b} - {today.day} {today:%b %Y}"


def period_from_dates(date_strings):
    """
    Build a period label from a list/Series of 'dd/mm/yyyy' strings.
    Falls back to default_period_text() if no valid dates.
    """
    ds = pd.to_datetime(pd.Series(date_strings), format="%d/%m/%Y", errors="coerce").dropna()
    if ds.empty:
        return default_period_text()
    lo, hi = ds.min(), ds.max()
    if (lo.year, lo.month) == (hi.year, hi.month):
        return f"{lo.day} {lo:%b} - {hi.day} {hi:%b %Y}"
    return f"{lo.day} {lo:%b %Y} - {hi.day} {hi:%b %Y}"


DEFAULT_PERIOD = default_period_text()

ROWS_PER_PAGE  = 38
BLOCK_OVERHEAD = 6
XL_THRESHOLD   = 20

# ---------- statement styling ----------
FONT_TITLE   = Font(name="Calibri", size=14, bold=True)
FONT_REF     = Font(name="Calibri", size=12, bold=True)
FONT_PERIOD  = Font(name="Calibri", size=10, italic=True)
FONT_HEADER  = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_BODY    = Font(name="Calibri", size=10)
FONT_BODY_SM = Font(name="Calibri", size=8)
FONT_TOTAL   = Font(name="Calibri", size=10, bold=True)
FONT_THANKS  = Font(name="Calibri", size=9, italic=True)

FILL_HEADER  = PatternFill("solid", fgColor="305496")
FILL_TOTAL   = PatternFill("solid", fgColor="D9E1F2")
FILL_REF     = PatternFill("solid", fgColor="FCE4D6")

THIN   = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")

HEADERS    = ["S.No", "Date", "Patient Name", "Investigation Done",
              "Investigation Charge", "Ambulance", "Discount",
              "Percent Cut", "Rate"]
COL_WIDTHS = [6, 12, 18, 30, 12, 11, 10, 11, 10]
LONG_TEST_CHARS = 25

# ---------- referrer normalization ----------
_SPACE_RE = re.compile(r"\s+")
_PUNCT_RE = re.compile(r"[.,]+$")


def normalize_referrer(name):
    s = str(name).strip()
    s = _SPACE_RE.sub(" ", s)
    s = _PUNCT_RE.sub("", s)
    return s.upper()


def is_html_bytes(b):
    head = b[:1024].lstrip().lower()
    return head.startswith(b"<") or b"<html" in head or b"<table" in head


def clean_money(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v))
    try:
        return float(s) if s not in ("", "-", ".") else 0.0
    except ValueError:
        return 0.0


def parse_date(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    try:
        return pd.to_datetime(v, dayfirst=True)
    except Exception:
        return None


def compute_rate(cut, disc, amb):
    a = amb if amb == 100 else 0
    return max(cut - disc - a, 0.0)


def load_source(raw_bytes, filename="file.xls"):
    if isinstance(raw_bytes, str):
        raw_bytes = raw_bytes.encode("utf-8", errors="ignore")

    if is_html_bytes(raw_bytes):
        try:
            tables = pd.read_html(io.BytesIO(raw_bytes))
        except Exception as e:
            raise ValueError(f"Could not parse HTML table: {e}")
        if not tables:
            raise ValueError("No <table> found in the HTML file.")
        df = max(tables, key=lambda t: t.shape[0] * t.shape[1])
    else:
        engine = "xlrd" if filename.lower().endswith(".xls") else "openpyxl"
        try:
            df = pd.read_excel(io.BytesIO(raw_bytes), engine=engine)
        except Exception:
            alt = "openpyxl" if engine == "xlrd" else "xlrd"
            df = pd.read_excel(io.BytesIO(raw_bytes), engine=alt)

    df.columns = [str(c).strip() for c in df.columns]

    aliases = {
        "RefNo":       ["RefNo", "Ref No", "Reference"],
        "PatientName": ["PatientName", "Patient Name"],
        "BillDate":    ["BillDate", "Bill Date", "Date"],
        "ReferBy":     ["ReferBy", "Refer By", "Referrer"],
        "TestName":    ["TestName", "Test Name", "Investigation"],
        "PatientRate": ["PatientRate", "Patient Rate", "Charge",
                        "Investigation Charge"],
        "DiscPercent": ["DiscPercent", "Disc", "Discount"],
        "CutRate":     ["CutRate", "PercentCut", "Percent Cut"],
        "Ambulance":   ["Ambulance", "Ambulance Charge", "Amb", "AMB",
                        "Ambulance<br>", "Ambulance\n", "Ambulance "],
    }
    rename = {}
    for canonical, options in aliases.items():
        for opt in options:
            if opt in df.columns and opt != canonical:
                rename[opt] = canonical
                break
    if rename:
        df = df.rename(columns=rename)

    required = ["PatientName", "BillDate", "ReferBy", "TestName",
                "PatientRate", "DiscPercent", "CutRate", "Ambulance"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"Missing columns: {missing}\nFound: {list(df.columns)}")

    for c in ["Rate", "Profit"]:
        if c in df.columns:
            df = df.drop(columns=[c])

    referby_str = df["ReferBy"].astype(str).str.strip().str.upper()
    df = df[~referby_str.isin(["TOTAL", "GRAND TOTAL", "SUM", ""])]

    if "RefNo" in df.columns:
        refno_str = df["RefNo"].astype(str).str.strip()
        df = df[~refno_str.isin(["", "nan", "TOTAL"])]

    df["ReferBy"] = df["ReferBy"].astype(str).apply(normalize_referrer)

    for c in ["PatientRate", "DiscPercent", "CutRate", "Ambulance"]:
        df[c] = df[c].apply(clean_money)

    df["PatientName"] = df["PatientName"].astype(str).str.strip()
    df["TestName"]    = df["TestName"].astype(str).str.strip()

    df["_date"] = df["BillDate"].apply(parse_date)
    df = df.sort_values(by=["ReferBy", "_date"], kind="stable").reset_index(drop=True)
    df["BillDate"] = df["_date"].apply(
        lambda d: d.strftime("%d/%m/%Y") if d is not None else ""
    )
    df = df.drop(columns=["_date"])

    df = df[df["ReferBy"].notna() & (df["ReferBy"] != "") & (df["ReferBy"] != "nan")]

    df["Rate"] = df.apply(
        lambda r: compute_rate(r["CutRate"], r["DiscPercent"], r["Ambulance"]),
        axis=1,
    )
    return df.reset_index(drop=True)


def fetch_from_gdrive(url_or_id):
    s = url_or_id.strip()
    if not s:
        raise ValueError("Empty link.")

    fid = None
    for pat in (r"/file/d/([A-Za-z0-9_-]+)",
                r"/spreadsheets/d/([A-Za-z0-9_-]+)",
                r"[?&]id=([A-Za-z0-9_-]+)"):
        m = re.search(pat, s)
        if m:
            fid = m.group(1)
            break
    if not fid:
        fid = s

    if "docs.google.com/spreadsheets" in s or "/spreadsheets/" in s:
        export_url = f"https://docs.google.com/spreadsheets/d/{fid}/export?format=xlsx"
    else:
        export_url = f"https://drive.google.com/uc?export=download&id={fid}"

    req = urllib.request.Request(export_url,
                                 headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            data = r.read()
    except Exception as e:
        raise ValueError(f"Fetch failed: {e}")

    if not data or len(data) < 100:
        raise ValueError("Downloaded file is empty or too small. "
                         "Make sure the file is shared as 'Anyone with the link'.")
    return data


def pack_referrers(ref_groups):
    sheets = []
    small, medium, large, xl = [], [], [], []

    for name, g in ref_groups:
        n = len(g)
        if n >= XL_THRESHOLD:
            xl.append((name, g))
        elif n >= 10:
            large.append((name, g))
        elif n >= 3:
            medium.append((name, g))
        else:
            small.append((name, g))

    for i, (name, g) in enumerate(xl, 1):
        sheets.append((f"XL{i}", [(name, g)]))

    def fill(bucket, cap, prefix):
        out, cur, used = [], [], 0
        for name, g in bucket:
            need = BLOCK_OVERHEAD + len(g)
            if used + need > cap and cur:
                out.append(cur)
                cur, used = [], 0
            cur.append((name, g))
            used += need
        if cur:
            out.append(cur)
        for i, grp in enumerate(out, 1):
            sheets.append((f"{prefix}{i}", grp))

    fill(large,  ROWS_PER_PAGE, "L")
    fill(medium, ROWS_PER_PAGE, "M")
    fill(small,  ROWS_PER_PAGE, "P")
    return sheets


def write_index(ws, ref_groups, sheet_map, period_text):
    FONT_IDX_TITLE  = Font(name="Calibri", size=14, bold=True)
    FONT_IDX_HEADER = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    FONT_IDX_BODY   = Font(name="Calibri", size=9)
    FONT_IDX_TOTAL  = Font(name="Calibri", size=9, bold=True)
    FONT_IDX_LINK   = Font(name="Calibri", size=9, color="0563C1", underline="single")

    IDX_HEADERS = ["Referrer", "Lines", "Rate (₹)",
                   "S/A", "F/NF", "D/ND", "ONLINE", "Go"]

    ws["A1"] = f"Commission Report Index — {period_text}"
    ws["A1"].font = FONT_IDX_TITLE
    ws.merge_cells("A1:Q1")
    ws.row_dimensions[1].height = 20

    table_starts = [1, 10]

    block_widths = [26, 7, 11, 6, 7, 7, 9, 5]
    col_widths = {}
    for start in table_starts:
        for off, w in enumerate(block_widths):
            col_widths[start + off] = w
    col_widths[9] = 2
    for c, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    n = len(ref_groups)
    per_col = (n + 1) // 2
    columns = [ref_groups[0:per_col],
               ref_groups[per_col:2*per_col]]

    header_row = 2
    start_row  = 3

    grand_total = 0.0
    grand_items = 0

    for start_col, group in zip(table_starts, columns):
        for i, h in enumerate(IDX_HEADERS):
            c = ws.cell(row=header_row, column=start_col + i, value=h)
            c.font = FONT_IDX_HEADER
            c.fill = FILL_HEADER
            c.alignment = CENTER
            c.border = BORDER
        ws.row_dimensions[header_row].height = 14

        r = start_row
        for name, g in group:
            sn, fr = sheet_map[name]

            c = ws.cell(row=r, column=start_col, value=name)
            c.font = FONT_IDX_BODY; c.border = BORDER; c.alignment = LEFT

            c = ws.cell(row=r, column=start_col + 1, value=len(g))
            c.font = FONT_IDX_BODY; c.border = BORDER; c.alignment = CENTER

            c = ws.cell(row=r, column=start_col + 2, value=float(g["Rate"].sum()))
            c.font = FONT_IDX_BODY; c.border = BORDER; c.alignment = RIGHT
            c.number_format = "#,##0"

            for i in range(3, 7):
                c = ws.cell(row=r, column=start_col + i, value=None)
                c.font = FONT_IDX_BODY; c.border = BORDER; c.alignment = CENTER

            link = ws.cell(row=r, column=start_col + 7, value="►")
            link.hyperlink = f"#'{sn}'!A{fr}"
            link.font = FONT_IDX_LINK; link.border = BORDER; link.alignment = CENTER

            ws.row_dimensions[r].height = 11
            grand_total += float(g["Rate"].sum())
            grand_items += len(g)
            r += 1

        c = ws.cell(row=r, column=start_col, value="SUBTOTAL")
        c.font = FONT_IDX_TOTAL; c.fill = FILL_TOTAL
        c.border = BORDER; c.alignment = LEFT

        c = ws.cell(row=r, column=start_col + 1,
                    value=sum(len(g) for _, g in group))
        c.font = FONT_IDX_TOTAL; c.fill = FILL_TOTAL
        c.border = BORDER; c.alignment = CENTER

        c = ws.cell(row=r, column=start_col + 2,
                    value=sum(float(g["Rate"].sum()) for _, g in group))
        c.font = FONT_IDX_TOTAL; c.fill = FILL_TOTAL
        c.border = BORDER; c.alignment = RIGHT
        c.number_format = "#,##0"

        for i in range(3, 8):
            c = ws.cell(row=r, column=start_col + i)
            c.fill = FILL_TOTAL; c.border = BORDER

    last_row = start_row + per_col + 1
    ws.cell(row=last_row + 1, column=1, value="GRAND TOTAL")
    ws.cell(row=last_row + 1, column=2, value=grand_items)
    c = ws.cell(row=last_row + 1, column=3, value=grand_total)
    c.number_format = "#,##0"
    for j in range(1, 9):
        cc = ws.cell(row=last_row + 1, column=j)
        cc.font = FONT_IDX_TOTAL
        cc.fill = FILL_TOTAL
        cc.border = BORDER
        cc.alignment = CENTER if j in (2, 4, 5, 6, 7, 8) else (LEFT if j == 1 else RIGHT)
    ws.row_dimensions[last_row + 1].height = 14

    ws.freeze_panes = "A3"

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:Q{last_row + 1}"


def write_block(ws, start_row, referrer, df, period_text):
    r = start_row

    ws.cell(row=r, column=1, value=referrer)
    ws.cell(row=r, column=1).font = FONT_REF
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(row=r, column=7, value="Thanks & Regards")
    ws.cell(row=r, column=7).font = FONT_THANKS
    ws.cell(row=r, column=7).alignment = RIGHT
    ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9)
    for c in range(1, 10):
        ws.cell(row=r, column=c).fill = FILL_REF
        ws.cell(row=r, column=c).border = BORDER
    ws.row_dimensions[r].height = 14
    r += 1

    ws.cell(row=r, column=1, value=period_text)
    ws.cell(row=r, column=1).font = FONT_PERIOD
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    ws.row_dimensions[r].height = 12
    r += 1

    header_row = r
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[r].height = 14
    r += 1

    for i, (_, row) in enumerate(df.iterrows(), 1):
        amb = row["Ambulance"]
        amb_disp = -100 if amb and amb != 0 else ""
        test_name = str(row["TestName"])
        test_font = FONT_BODY_SM if len(test_name) > LONG_TEST_CHARS else FONT_BODY

        vals = [
            i,
            row["BillDate"],
            row["PatientName"],
            test_name,
            row["PatientRate"],
            amb_disp,
            row["DiscPercent"],
            row["CutRate"],
            row["Rate"],
        ]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = test_font if j == 4 else FONT_BODY
            c.border = BORDER
            c.alignment = (
                CENTER if j in (1, 2, 6)
                else RIGHT if j in (5, 7, 8, 9)
                else LEFT
            )
            if j in (5, 6, 7, 8, 9):
                c.number_format = "#,##0"
        r += 1

    ws.cell(row=r, column=1, value="Total")
    ws.cell(row=r, column=9, value=float(df["Rate"].sum()))
    for j in range(1, 10):
        c = ws.cell(row=r, column=j)
        c.font = FONT_TOTAL
        c.fill = FILL_TOTAL
        c.border = BORDER
        c.alignment = CENTER if j == 1 else RIGHT
        if j == 9:
            c.number_format = "#,##0"
    r += 1

    ws.cell(row=r, column=1, value="Thanks & Regards")
    ws.cell(row=r, column=1).font = FONT_THANKS
    r += 1

    return r, header_row


def build_workbook(df, ref_groups, packed, period_text):
    wb = Workbook()
    wb.remove(wb.active)
    idx = wb.create_sheet("Index")
    sheet_map = {}

    for sn, group in packed:
        ws = wb.create_sheet(sn)
        for i, w in enumerate(COL_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize   = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth  = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left   = 0.3
        ws.page_margins.right  = 0.3
        ws.page_margins.top    = 0.4
        ws.page_margins.bottom = 0.4

        if len(group) == 1:
            ws.print_title_rows = "1:3"

        row = 1
        for referrer, g in group:
            row, hr = write_block(ws, row, referrer, g, period_text)
            sheet_map[referrer] = (sn, hr - 2)

    write_index(idx, ref_groups, sheet_map, period_text)
    return wb


def build_individual_workbook(referrer, df_one, period_text):
    wb = Workbook()
    ws = wb.active
    safe_name = re.sub(r"[\[\]:*?/\\]", "-", str(referrer))[:31]
    ws.title = safe_name or "Statement"

    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize   = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth  = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left   = 0.3
    ws.page_margins.right  = 0.3
    ws.page_margins.top    = 0.4
    ws.page_margins.bottom = 0.4
    ws.print_title_rows = "1:3"

    write_block(ws, 1, referrer, df_one, period_text)
    return wb


# ============================================================
#                     STREAMLIT UI
# ============================================================

st.set_page_config(page_title="Commission Report Builder",
                   page_icon="📊", layout="centered")

st.title("📊 Commission Report Builder")
st.caption("Load your flat Excel file (.xls / .xlsx / HTML-exported .xls) "
           "from an upload or a Google Drive / Sheets link.")

if "raw_bytes"   not in st.session_state: st.session_state.raw_bytes   = None
if "source_name" not in st.session_state: st.session_state.source_name = "source.xls"
if "report_buf"  not in st.session_state: st.session_state.report_buf  = None
if "summary"     not in st.session_state: st.session_state.summary     = None
if "df_cache"    not in st.session_state: st.session_state.df_cache    = None

if "period_text_input" not in st.session_state:
    st.session_state["period_text_input"] = default_period_text()

period_text = st.text_input("Report period", key="period_text_input")

st.subheader("Option 1 — Upload")
uploaded = st.file_uploader("Choose your Excel file",
                            type=["xls", "xlsx", "html", "htm"],
                            key="file_uploader_widget")
if uploaded is not None:
    new_bytes = uploaded.read()
    if new_bytes != st.session_state.raw_bytes:
        st.session_state.raw_bytes   = new_bytes
        st.session_state.source_name = uploaded.name
        st.session_state.report_buf  = None
        st.session_state.df_cache    = None
        st.session_state.pop("_referrer_dataset_id", None)

st.subheader("Option 2 — Google Drive / Sheets link")
st.caption("Share the file as 'Anyone with the link', then paste the link here.")
drive_input = st.text_input("Drive / Sheets link or file ID", value="",
                            key="drive_link_input")

if st.button("Fetch from Drive", key="fetch_drive_btn"):
    if not drive_input.strip():
        st.error("Paste a link first.")
    else:
        with st.spinner("Downloading from Google…"):
            try:
                st.session_state.raw_bytes   = fetch_from_gdrive(drive_input)
                st.session_state.source_name = "drive_file.xls"
                st.session_state.report_buf  = None
                st.session_state.df_cache    = None
                st.session_state.pop("_referrer_dataset_id", None)
                st.success("File fetched successfully.")
            except Exception as e:
                st.error(f"Drive fetch failed: {e}")

if st.session_state.raw_bytes is not None:
    st.info(f"Loaded file: {st.session_state.source_name} "
            f"({len(st.session_state.raw_bytes):,} bytes)")
    if st.button("Generate Report", type="primary", key="generate_report_btn"):
        with st.spinner("Building report…"):
            try:
                raw_bytes = st.session_state.raw_bytes
                if is_html_bytes(raw_bytes):
                    tables = pd.read_html(io.BytesIO(raw_bytes))
                    raw_df = max(tables, key=lambda t: t.shape[0] * t.shape[1])
                else:
                    engine = "xlrd" if st.session_state.source_name.lower().endswith(".xls") else "openpyxl"
                    try:
                        raw_df = pd.read_excel(io.BytesIO(raw_bytes), engine=engine)
                    except Exception:
                        alt = "openpyxl" if engine == "xlrd" else "xlrd"
                        raw_df = pd.read_excel(io.BytesIO(raw_bytes), engine=alt)

                raw_df.columns = [str(c).strip() for c in raw_df.columns]
                df = load_source(raw_bytes, st.session_state.source_name)
            except Exception as e:
                st.error(f"Could not read file: {e}")
                st.stop()

            if df.empty:
                st.warning("No rows found after cleaning. Check your file.")
                st.stop()

            # Auto-update period text from actual data dates
            data_period = period_from_dates(df["BillDate"])
            if st.session_state.get("period_text_input") in ("", DEFAULT_PERIOD):
                st.session_state["period_text_input"] = data_period
            period_text = st.session_state["period_text_input"]

            ref_groups = sorted(
                df.groupby("ReferBy", sort=False),
                key=lambda x: x[1]["Rate"].sum(),
                reverse=True,
            )
            packed = pack_referrers(ref_groups)

            packed_names = {n for _, group in packed for n, _ in group}
            expected = {n for n, _ in ref_groups}
            if packed_names != expected:
                st.error("Internal error: some referrers were dropped during packing.")
                st.stop()

            wb = build_workbook(df, ref_groups, packed, period_text)
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)

            computed_total = float(df["Rate"].sum())

            # Anomaly detection: compare source Rate column to our calculated Rate
            source_rate_col = next(
                (c for c in raw_df.columns if str(c).strip().upper() == "RATE"),
                None,
            )
            anomalies = []
            if source_rate_col is not None:
                clean_raw_rate = raw_df[source_rate_col].apply(clean_money)
                clean_calc_rate = df["Rate"]
                min_len = min(len(clean_raw_rate), len(clean_calc_rate))
                diff = (
                    clean_raw_rate.iloc[:min_len].reset_index(drop=True)
                    - clean_calc_rate.iloc[:min_len].reset_index(drop=True)
                )
                diff_mask = diff.abs() > 0.01
                if diff_mask.any():
                    anomaly_df = df.iloc[:min_len][diff_mask].copy().reset_index(drop=True)
                    anomaly_df["File_Rate"] = clean_raw_rate.iloc[:min_len][diff_mask].values
                    anomaly_df["App_Rate"]  = clean_calc_rate.iloc[:min_len][diff_mask].values
                    anomalies = anomaly_df[
                        ["BillDate", "PatientName", "TestName",
                         "CutRate", "File_Rate", "App_Rate"]
                    ].to_dict("records")

            st.session_state.report_buf = buf.getvalue()
            st.session_state.df_cache   = df
            st.session_state.summary = {
                "rows": int(len(df)),
                "referrers": int(df["ReferBy"].nunique()),
                "total_rate": computed_total,
                "period": period_text,
                "anomalies": anomalies,
                "top": [
                    (n, len(g), float(g["Rate"].sum()))
                    for n, g in ref_groups[:15]
                ],
            }
            st.session_state.pop("_referrer_dataset_id", None)

if st.session_state.report_buf is not None:
    s = st.session_state.summary
    st.success(f"✅ Loaded **{s['rows']}** rows across **{s['referrers']}** referrers "
               f"— period **{s.get('period', period_text)}**.")
    st.metric("Computed total Rate (₹)", f"{s['total_rate']:,.0f}")

    if s.get("anomalies"):
        st.error(f"⚠️ **Calculation Anomaly Detected!** The app found "
                 f"**{len(s['anomalies'])}** rows where the source file's "
                 f"Rate differs from our calculated Rate.")
        with st.expander("🔍 View Anomaly Rows", expanded=True):
            st.dataframe(pd.DataFrame(s["anomalies"]),
                         use_container_width=True, hide_index=True)

    st.download_button(
        label="⬇️ Download combined report (.xlsx)",
        data=st.session_state.report_buf,
        file_name="Commission_Report_Output.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="download_combined_report",
    )

    with st.expander("Preview referrers (top 15 by total Rate)"):
        preview = pd.DataFrame({
            "Referrer": [r for r, _, _ in s["top"]],
            "Line Items": [n for _, n, _ in s["top"]],
            "Total Rate (₹)": [f"{t:,.0f}" for _, _, t in s["top"]],
        })
        st.dataframe(preview, hide_index=True, use_container_width=True)

if st.session_state.df_cache is not None:
    st.divider()
    with st.expander("🚑 Ambulance Audit — verify what's in the source file",
                     expanded=False):
        df_amb = st.session_state.df_cache

        st.write("**Ambulance column sanity check**")
        st.write("Columns found:", list(df_amb.columns))
        st.write("Ambulance dtype:", str(df_amb["Ambulance"].dtype))
        st.write("Unique ambulance values (up to 30):",
                 sorted(df_amb["Ambulance"].unique().tolist())[:30])
        st.write("Value counts:",
                 df_amb["Ambulance"].value_counts().to_dict())

        amb_df = df_amb[df_amb["Ambulance"] > 0]
        st.write(f"**Total rows with Ambulance > 0: {len(amb_df)}**")

        if amb_df.empty:
            st.warning("No rows have a non-zero Ambulance value. "
                       "If you expect ambulance charges, the column may not "
                       "be parsing correctly — compare the columns listed above "
                       "against your source file.")
        else:
            st.write("**Referrers with at least one ambulance charge:**")
            per_ref = (
                amb_df.groupby("ReferBy")
                      .agg(Ambulance_Rows=("Ambulance", "size"),
                           Total_Ambulance=("Ambulance", "sum"))
                      .sort_values("Total_Ambulance", ascending=False)
                      .reset_index()
                      .rename(columns={"ReferBy": "Referrer"})
            )
            st.dataframe(per_ref, use_container_width=True, hide_index=True)

            st.write("**Every ambulance row in the file:**")
            show_cols = [c for c in
                         ["RefNo", "PatientName", "ReferBy", "BillDate",
                          "TestName", "Ambulance"]
                         if c in amb_df.columns]
            st.dataframe(
                amb_df[show_cols].reset_index(drop=True),
                use_container_width=True,
                hide_index=True,
            )

if st.session_state.df_cache is not None:
    st.divider()
    st.subheader("Individual Referrer Statement")
    st.caption("Pick a referrer to download a single-sheet XLSX containing only "
               "that referrer's statement, in the same layout as the combined report.")

    df = st.session_state.df_cache
    referrer_list = sorted(df["ReferBy"].unique().tolist())

    if not referrer_list:
        st.warning("No referrers available in the loaded data.")
    else:
        dataset_id = (
            st.session_state.get("source_name", ""),
            int(len(df)),
        )

        if st.session_state.get("_referrer_dataset_id") != dataset_id:
            st.session_state.pop("referrer_pick", None)
            st.session_state.pop("referrer_search_query", None)
            st.session_state["_referrer_dataset_id"] = dataset_id

        query = st.text_input(
            "Search referrer (optional)",
            key="referrer_search_query",
        ).strip().upper()

        filtered = [r for r in referrer_list if query in r] if query else referrer_list

        if not filtered:
            st.warning("No referrer matches your search.")
        else:
            current = st.session_state.get("referrer_pick")
            idx = filtered.index(current) if
